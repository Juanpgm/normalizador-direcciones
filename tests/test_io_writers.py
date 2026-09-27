"""Writer registry: streaming sinks, atomic output, failure modes."""

from __future__ import annotations

import json
import os

import openpyxl
import pandas as pd
import pytest

from dataset_helpers import assert_same_frames

from cali_address.io import (
    SinkWriteError,
    UnsupportedFormatError,
    infer_sink_format,
    list_writer_formats,
    open_sink,
)

COLS = ["id", "direccion", "estado", "lat", "lon", "confianza"]


def _chunks():
    a = pd.DataFrame({"id": ["1", "2"], "direccion": ["CL 5 # 38 - 20", "Carrera 100\n# 15-30, \"x\""],
                      "estado": ["OK", "SIN_MATCH"], "lat": [3.4, None], "lon": [-76.5, None],
                      "confianza": [0.9, None]})
    b = pd.DataFrame({"id": ["3"], "direccion": ["Ñandú"], "estado": ["OK"], "lat": [3.5],
                      "lon": [-76.6], "confianza": [0.8]})
    return [a, b]


def _write(path, chunks, **kw):
    with open_sink(str(path), **kw) as sink:
        for c in chunks:
            sink.write(c)


def test_writer_formats_listed():
    assert {"csv", "tsv", "xlsx", "parquet", "json", "jsonl", "geojson"} <= set(list_writer_formats())


@pytest.mark.parametrize("name,fmt", [("a.csv", "csv"), ("a.tsv", "tsv"), ("a.xlsx", "xlsx"),
                                       ("a.parquet", "parquet"), ("a.json", "json"),
                                       ("a.jsonl", "jsonl"), ("a.geojson", "geojson")])
def test_infer_sink_format(name, fmt):
    assert infer_sink_format(name) == fmt


def test_unsupported_output_extension(tmp_path):
    with pytest.raises(UnsupportedFormatError):
        open_sink(str(tmp_path / "out.docx"))
    with pytest.raises(UnsupportedFormatError):
        open_sink(str(tmp_path / "out"))


def test_csv_round_trip_with_bom_and_quotes(tmp_path):
    p = tmp_path / "o.csv"
    _write(p, _chunks())
    assert p.read_bytes().startswith(b"\xef\xbb\xbf")  # utf-8-sig, same as the legacy writer
    back = pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    assert back["direccion"].tolist() == ["CL 5 # 38 - 20", "Carrera 100\n# 15-30, \"x\"", "Ñandú"]
    assert len(back) == 3


def test_tsv(tmp_path):
    p = tmp_path / "o.tsv"
    _write(p, _chunks())
    assert len(pd.read_csv(p, sep="\t", encoding="utf-8-sig")) == 3


def test_xlsx_has_three_sheets(tmp_path):
    p = tmp_path / "o.xlsx"
    _write(p, _chunks())
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["normalizado", "revisar", "resumen"]
    assert wb["normalizado"].max_row == 4
    wb.close()


def test_parquet_first_chunk_all_null_then_values(tmp_path):
    a = pd.DataFrame({"id": ["1"], "lat": [None], "nota": [None], "estado": ["SIN_MATCH"]})
    b = pd.DataFrame({"id": ["2"], "lat": [3.4], "nota": ["algo"], "estado": ["OK"]})
    p = tmp_path / "o.parquet"
    _write(p, [a, b])
    back = pd.read_parquet(p)
    assert len(back) == 2 and back.loc[1, "nota"] == "algo" and back.loc[1, "lat"] == 3.4


def test_parquet_mixed_types_in_object_column(tmp_path):
    a = pd.DataFrame({"v": [1, 2]}, dtype=object)
    b = pd.DataFrame({"v": ["x", 3.5]}, dtype=object)
    p = tmp_path / "o.parquet"
    _write(p, [a, b])
    assert len(pd.read_parquet(p)) == 4


def test_jsonl_and_json(tmp_path):
    pj = tmp_path / "o.jsonl"
    _write(pj, _chunks())
    lines = pj.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 3 and json.loads(lines[1])["lat"] is None
    pa = tmp_path / "o.json"
    _write(pa, _chunks())
    data = json.loads(pa.read_text(encoding="utf-8"))
    assert len(data) == 3 and data[2]["direccion"] == "Ñandú"


def test_json_zero_rows_is_empty_array(tmp_path):
    p = tmp_path / "o.json"
    with open_sink(str(p)) as sink:
        pass
    assert json.loads(p.read_text(encoding="utf-8")) == []


def test_geojson_keeps_feature_count_and_null_geometry(tmp_path):
    p = tmp_path / "o.geojson"
    _write(p, _chunks())
    fc = json.loads(p.read_text(encoding="utf-8"))
    assert fc["type"] == "FeatureCollection" and len(fc["features"]) == 3
    assert fc["features"][1]["geometry"] is None
    assert fc["features"][0]["geometry"]["coordinates"] == [-76.5, 3.4]


def test_non_json_native_values_are_serialised(tmp_path):
    df = pd.DataFrame({"id": ["1"], "fecha": [pd.Timestamp("2024-01-02")], "n": [pd.Series([1]).iloc[0]]})
    p = tmp_path / "o.jsonl"
    _write(p, [df])
    row = json.loads(p.read_text(encoding="utf-8").splitlines()[0])
    assert row["fecha"].startswith("2024-01-02") and row["n"] == 1


def test_header_only_output_from_empty_frame(tmp_path):
    p = tmp_path / "o.csv"
    _write(p, [pd.DataFrame(columns=COLS)])
    assert pd.read_csv(p, encoding="utf-8-sig").columns.tolist() == COLS


def test_creates_missing_parent_directories(tmp_path):
    p = tmp_path / "a" / "b" / "o.csv"
    _write(p, _chunks())
    assert p.exists()


def test_output_path_is_a_directory(tmp_path):
    d = tmp_path / "out.csv"
    d.mkdir()
    with pytest.raises(SinkWriteError):
        open_sink(str(d))


def test_output_parent_is_a_file(tmp_path):
    f = tmp_path / "file.txt"
    f.write_text("x")
    with pytest.raises(SinkWriteError):
        open_sink(str(f / "o.csv"))


def test_failure_leaves_no_partial_file_and_keeps_previous_output(tmp_path):
    p = tmp_path / "o.csv"
    p.write_text("previous,run\n1,2\n", encoding="utf-8")
    with pytest.raises(RuntimeError):
        with open_sink(str(p)) as sink:
            sink.write(_chunks()[0])
            raise RuntimeError("boom")
    assert p.read_text(encoding="utf-8") == "previous,run\n1,2\n"
    assert [f for f in os.listdir(tmp_path) if f != "o.csv"] == []


def test_success_replaces_previous_output(tmp_path):
    p = tmp_path / "o.csv"
    p.write_text("previous\n", encoding="utf-8")
    _write(p, _chunks())
    assert "previous" not in p.read_text(encoding="utf-8-sig")


def test_memory_sink_collects_frames():
    from cali_address.io import MemorySink

    sink = MemorySink()
    for c in _chunks():
        sink.write(c)
    sink.close()
    assert_same_frames(sink.frame, pd.concat(_chunks(), ignore_index=True))
